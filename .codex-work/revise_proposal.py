from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(r"C:\Users\HP\Downloads\Final year project")
SOURCE = ROOT / "not_useful" / "documents" / "expanded_final_defence" / "Final YEAR PROPOSAL.docx"
OUTPUT = ROOT / "not_useful" / "documents" / "expanded_final_defence" / "Final YEAR PROPOSAL Revised.docx"


def norm(text: str) -> str:
    return " ".join((text or "").split())


def para_index(doc: Document, text: str) -> int:
    target = norm(text)
    for i, p in enumerate(doc.paragraphs):
        if norm(p.text) == target:
            return i
    for i, p in enumerate(doc.paragraphs):
        if norm(p.text).startswith(target):
            return i
    raise ValueError(f"Paragraph not found: {text}")


def find_first(doc: Document, startswith: str, start: int = 0) -> int:
    for i, p in enumerate(doc.paragraphs[start:], start):
        if norm(p.text).startswith(startswith):
            return i
    raise ValueError(f"Paragraph starting with not found: {startswith}")


def set_text(paragraph, text: str) -> None:
    paragraph.text = text
    paragraph.paragraph_format.space_after = Pt(6)


def insert_paragraph_before(doc: Document, before_text: str, text: str, style: str = "Body Text"):
    target = doc.paragraphs[para_index(doc, before_text)]
    p = target.insert_paragraph_before(text, style=style)
    p.paragraph_format.space_after = Pt(6)
    return p


def insert_many_before(doc: Document, before_text: str, items: list[tuple[str, str]]):
    target = doc.paragraphs[para_index(doc, before_text)]
    for text, style in items:
        p = target.insert_paragraph_before(text, style=style)
        p.paragraph_format.space_after = Pt(6)


def insert_table_before(doc: Document, before_text: str, rows: list[list[str]], style: str = "Table Grid"):
    target = doc.paragraphs[para_index(doc, before_text)]
    table = doc.add_table(rows=0, cols=len(rows[0]))
    try:
        table.style = style
    except Exception:
        pass
    for row_data in rows:
        row = table.add_row()
        for cell, value in zip(row.cells, row_data):
            cell.text = value
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(9)
    target._p.addprevious(table._tbl)
    return table


def add_reference(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text, style="Body Text")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.5)
    p.paragraph_format.space_after = Pt(6)


def replace_everywhere(doc: Document, replacements: dict[str, str]) -> None:
    for p in doc.paragraphs:
        text = p.text
        updated = text
        for old, new in replacements.items():
            updated = updated.replace(old, new)
        if updated != text:
            set_text(p, updated)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    text = p.text
                    updated = text
                    for old, new in replacements.items():
                        updated = updated.replace(old, new)
                    if updated != text:
                        set_text(p, updated)


def append_objective_row(table):
    existing = [
        (
            "1.",
            "To identify and characterize common failure patterns in data pipeline batch jobs.",
            "Collect pipeline execution logs, workload traces, and ETL demonstration evidence; extract severity, keyword, retry, resource, timing, and task-failure features.",
        ),
        (
            "2.",
            "To carry out a comparative analysis based on accuracy metrics for multiple machine learning algorithms to effectively predict pipeline failures.",
            "Compare XGBoost with logistic regression, decision tree, random forest, histogram gradient boosting, and a majority baseline using saved evaluation outputs.",
        ),
        (
            "3.",
            "To develop a prediction model to continuously monitor pipeline execution of batch jobs.",
            "Train source-specific prediction artifacts and integrate the selected ETL/HDFS-compatible model with the live runner and predictor plugin.",
        ),
        (
            "4.",
            "To test and evaluate the system using Precision, Recall and F1-Score.",
            "Compute confusion-matrix counts and evaluate precision, recall, F1-score, ROC-AUC, PR-AUC, MCC, specificity, log loss, and inference latency.",
        ),
    ]
    for row, values in zip(table.rows[1:5], existing):
        for cell, value in zip(row.cells, values):
            cell.text = value
    row = table.add_row()
    row.cells[0].text = "5."
    row.cells[1].text = "To visually represent the monitoring workflow through a dashboard and orchestration view."
    row.cells[2].text = (
        "Develop a Streamlit dashboard and Airflow DAG evidence to present pipeline stages, "
        "logs, risk classifications, review controls, model comparisons, and generated reports."
    )


def main() -> None:
    doc = Document(str(SOURCE))

    replace_everywhere(
        doc,
        {
            "A PROJECT PROPOSAL SUBMITTED": "A PROJECT REPORT SUBMITTED",
            "LITERATUR REVIEW": "LITERATURE REVIEW",
            "factors likc transaction constancy": "factors like transaction consistency",
            "modelartifact/model-comparisons.md": "model artifact and model-comparisons.md",
            "ET L model": "ETL model",
            "Render. yaml": "render.yaml",
            "python-version": ".python-version",
            "(Streamlit, 2026.)": "(Streamlit, n.d.)",
            "(Apache Software Foundation, 2026)": "(Apache Software Foundation, n.d.)",
            "(scikit-learn developers, 2026)": "(scikit-learn developers, n.d.)",
            "Apache Software Foundation. (2026.).": "Apache Software Foundation. (n.d.).",
            "Streamlit. (2026.).": "Streamlit. (n.d.).",
            "scikit-learn developers. (2026). f1_score.": "scikit-learn developers. (n.d.-b). f1_score.",
            "scikit-learn developers. (2026). StratifiedGroupKFold.": "scikit-learn developers. (n.d.-a). StratifiedGroupKFold.",
        },
    )

    # Strengthen Chapter One.
    set_text(
        doc.paragraphs[find_first(doc, "The goal of this study is to apply XGBoost")],
        (
            "The goal of this study is to apply XGBoost to ETL and data-pipeline batch-job evidence "
            "so that likely failures can be identified before they produce downstream data-quality "
            "or reporting problems. The expected contribution is a defensible monitoring workflow "
            "that connects log ingestion, feature extraction, model prediction, risk classification, "
            "dashboard presentation, and report generation. This closing focus leads directly to the "
            "problem addressed in the next section: current monitoring practices often reveal failures "
            "only after they have already affected the pipeline."
        ),
    )
    set_text(
        doc.paragraphs[find_first(doc, "The non-observability of data pipeline batch job failures")],
        (
            "The non-observability of data pipeline batch job failures in enterprise settings can allow "
            "errors to remain unnoticed until they create production data-quality problems. Manual log "
            "analysis keeps data engineering teams in a reactive maintenance position instead of enabling "
            "proactive pipeline optimization (Akash & Charate, 2025). Recent data-pipeline quality research "
            "also shows that data type, cleaning, ingestion, and compatibility issues remain common sources "
            "of pipeline defects, which strengthens the need for earlier automated detection (Foidl et al., 2024)."
        ),
    )
    set_text(
        doc.paragraphs[find_first(doc, "These failure patterns are not automated")],
        (
            "Existing monitoring systems can often display logs, status messages, or completed failure alerts, "
            "but they do not always learn from historical failure patterns or convert live execution evidence "
            "into an early warning score. As a result, engineers may spend substantial time tracing integration "
            "and transformation issues after the problem has already affected reporting, analytics, or downstream "
            "business decisions (Chanda, 2024; Ogunsola et al., 2022). The gap addressed by this study is therefore "
            "the absence of an integrated, machine-learning-supported workflow that can characterize failure patterns, "
            "classify risk during execution, and present the evidence in a form that supports timely intervention."
        ),
    )

    set_text(
        doc.paragraphs[find_first(doc, "The study uses an agile methodology approach")],
        (
            "The study uses an agile methodology approach to develop the failure prediction application in small iterations "
            "that can continuously monitor pipeline execution for batch jobs. The system is evaluated using Precision, "
            "Recall, F1-Score and supporting metrics, and its monitoring workflow is represented visually through the "
            "Streamlit dashboard, generated reports, and Airflow DAG evidence."
        ),
    )

    # Add the missing fifth objective and align Table 1.1.
    methodology_heading = doc.paragraphs[find_first(doc, "Methodology")]
    methodology_heading.insert_paragraph_before(
        "to visually represent the monitoring workflow using a dashboard and orchestration evidence.",
        style="List Paragraph",
    )
    append_objective_row(doc.tables[2])

    set_text(
        doc.paragraphs[find_first(doc, "The evaluation tools were chosen in line")],
        (
            "The evaluation tools were chosen in line with the research aims. The F1-score was included because it "
            "combines the strengths of precision and recall and is a more appropriate measure than accuracy when "
            "failures are infrequent (scikit-learn developers, n.d.-b). Stratified and grouped validation was critical "
            "because related snapshots, captured from the same job, cannot be blindly treated as independent "
            "training/test data points (scikit-learn developers, n.d.-a)."
        ),
    )

    # Add directly comparable literature and summary comparison.
    insert_many_before(
        doc,
        "Gaps Identified in Literature",
        [
            (
                "Recent related work has also emphasized that log-anomaly detection must be evaluated from an operational "
                "perspective, not only from a model-score perspective. Ma et al. (2024) surveyed practitioners and reviewed "
                "log-anomaly detection studies, showing that adoption depends on whether tools meet practical needs such as "
                "trust, usability, and alignment with maintenance workflows. This is relevant to the present study because "
                "the Streamlit dashboard, review controls, and generated reports were designed to make prediction evidence "
                "usable to an operator rather than leaving it as a raw model output.",
                "Body Text",
            ),
            (
                "Wang et al. (2024) proposed a cross-system log-anomaly detection approach that addresses labeling cost, "
                "evolving logs, and adaptation across systems. Their study reinforces the importance of source-specific "
                "validation because a model trained in one environment may not automatically generalize to another. This "
                "supports the present project's decision to keep ETL/HDFS-compatible, Alibaba, and Google 2019 artifacts "
                "separate instead of forcing a single merged model across incompatible schemas.",
                "Body Text",
            ),
            (
                "Zhang et al. (2024) studied multivariate log-based anomaly detection for distributed databases and argued "
                "that single-node or single-source evidence may be insufficient in distributed settings. This finding is "
                "important for pipeline failure prediction because batch jobs may fail due to interactions among logs, "
                "resources, tasks, and dependencies. The present study responds to this by combining log features, workload "
                "features, resource context, dashboard evidence, and report generation.",
                "Body Text",
            ),
            ("Table 2.1: Summary comparison of selected related works", "Body Text"),
        ],
    )
    insert_table_before(
        doc,
        "Gaps Identified in Literature",
        [
            ["Study", "Dataset/context", "Method/focus", "Key finding", "Limitation/gap addressed by this project"],
            ["Jassas and Mahmoud (2022)", "Cloud-computing job records", "Machine-learning job failure prediction", "Showed that ML can support job-failure prediction in cloud workloads.", "Did not provide the same ETL-facing dashboard, report, and intervention workflow."],
            ["Du et al. (2017)", "System log sequences", "DeepLog with LSTM sequence modeling", "Demonstrated deep-learning-based log anomaly detection.", "Focused on sequence anomaly detection rather than an attachable ETL batch-job monitoring plugin."],
            ["Ma et al. (2024)", "Practitioner survey and literature review", "Adoption expectations for log anomaly detection", "Showed the importance of practical usability and trust.", "Supports the need for visual, reviewable prediction evidence in this project."],
            ["Wang et al. (2024)", "Multiple supercomputing log datasets", "Cross-system meta-learning for log anomalies", "Addressed adaptation across systems with limited labels.", "Reinforces the need for source-specific validation and cautious generalization."],
            ["Zhang et al. (2024)", "Distributed database logs", "Multivariate log-based anomaly detection", "Showed that multivariate evidence can improve anomaly detection.", "Supports combining logs, resource context, workload features, and reports."],
        ],
    )

    # Reproducibility detail in methodology.
    insert_many_before(
        doc,
        "Table 3.5: Stored model artifacts",
        [
            (
                "For reproducibility, the implemented XGBoost configuration used 300 estimators, a maximum tree depth of 5, "
                "a learning rate of 0.08, subsampling of 0.8, column sampling by tree of 0.8, and random_state set to 42. "
                "The train-test split reserved 25% of the available records for held-out testing. Where process or group "
                "identifiers were available, the implementation used a group-held-out split so that related snapshots from "
                "the same job were not divided across training and testing partitions.",
                "Body Text",
            ),
            (
                "The validation code also used five-fold cross-validation. StratifiedGroupKFold was applied when repeated "
                "snapshots or group identifiers made grouped validation appropriate; otherwise, stratified row-based folds "
                "were used. This explicit validation strategy addresses leakage risk and makes the model-development process "
                "more reproducible from the stored source code and artifacts (scikit-learn developers, n.d.-a).",
                "Body Text",
            ),
        ],
    )

    # Add Objective 1 results in Chapter Four.
    insert_many_before(
        doc,
        "The primary stored ETL/HDFS-compatible metrics achieve",
        [
            ("4.3.1 Failure Pattern and Feature-Importance Results", "Heading 3"),
            (
                "Objective 1 required the study to identify and characterize common failure patterns in data pipeline batch jobs. "
                "The saved feature-importance chart from the ETL/HDFS-compatible XGBoost model shows that the strongest warning "
                "signals were error_rate, info_count, error_count, max_error_burst, warn_count, and warn_rate. These features "
                "indicate that failure risk in the demonstration model was driven mainly by the concentration of error messages, "
                "the volume of normal and abnormal log activity, repeated error bursts, and warning density.",
                "Body Text",
            ),
            (
                "The result is operationally meaningful because the strongest features are observable while a pipeline is running. "
                "A high error rate or clustered burst of errors can alert the system before the final job outcome is known, while "
                "warning counts and warning rates can indicate degradation before failure. This closes the loop between the first "
                "objective, the feature-engineering method in Chapter Three, and the evaluation evidence presented in this chapter.",
                "Body Text",
            ),
            ("Table 4.2: Feature-importance and failure-pattern evidence from the ETL/HDFS-compatible model", "Body Text"),
        ],
    )
    insert_table_before(
        doc,
        "The primary stored ETL/HDFS-compatible metrics achieve",
        [
            ["Top signal/pattern", "Evidence from saved feature-importance chart", "Operational interpretation"],
            ["Error density", "error_rate and error_count ranked among the strongest predictors.", "Frequent errors indicate unstable execution and possible terminal failure."],
            ["Burst behavior", "max_error_burst appeared as a major feature.", "Repeated consecutive errors are more serious than isolated messages."],
            ["Warning activity", "warn_count and warn_rate were visible among the leading features.", "Warning accumulation can show degradation before a hard failure occurs."],
            ["Execution volume", "info_count and total/log-line indicators contributed to the model.", "The amount of execution evidence helps distinguish normal progress from abnormal runs."],
            ["Failure keywords and retries", "Keyword and retry features contributed at lower levels.", "Textual failure hints support prediction when combined with severity and burst features."],
        ],
    )
    replace_everywhere(
        doc,
        {
            "Table 4.2: Held-out XGBoost evaluation results preserved from project artifacts": "Table 4.3: Held-out XGBoost evaluation results preserved from project artifacts",
            "Table 4.3: Classifier comparison for the ETL/HDFS-compatible experiment": "Table 4.4: Classifier comparison for the ETL/HDFS-compatible experiment",
            "Table 4.4: Classifier comparison for bounded Alibaba and Google benchmark samples": "Table 4.5: Classifier comparison for bounded Alibaba and Google benchmark samples",
            "Table 4.5: Dashboard components and evidence value": "Table 4.6: Dashboard components and evidence value",
            "Table 4.6: Report payload evidence": "Table 4.7: Report payload evidence",
        },
    )

    # Strengthen Chapter Four discussion with explicit literature connection and neutral voice.
    insert_many_before(
        doc,
        "The overall result must therefore be interpreted cautiously.",
        [
            (
                "The ETL/HDFS-compatible F1-score of 0.9143 also compares favorably with the DeepLog study's reported "
                "top-1 prediction accuracy of 88.9%, although the metrics and datasets are not identical (Du et al., 2017). "
                "The comparison is useful because both studies show that log-derived evidence can reveal abnormal execution "
                "patterns, but the present project adds an ETL monitoring dashboard, report output, and review workflow.",
                "Body Text",
            ),
            (
                "The bounded Alibaba and Google results also align with the caution raised by job-failure prediction research: "
                "performance can vary by workload source, label imbalance, and feature schema (Jassas & Mahmoud, 2022). "
                "This explains why the report treats XGBoost as the selected implemented model while still recommending "
                "source-specific model comparison when the system is extended.",
                "Body Text",
            ),
        ],
    )
    replace_everywhere(
        doc,
        {
            "We use XGBoost as our predictor in this work": "This study uses XGBoost as the implemented predictor",
            "our model needs to be sourced-dependent": "the model needs to remain source-dependent",
            "Our results must be interpreted": "The results must be interpreted",
            "what this project aims to do, that is": "what this project aims to do; that is,",
            "our proposed end-to-end system": "the proposed end-to-end system",
            "we have not provided": "the study has not provided",
            "our result must be balanced": "the result must be balanced",
            "supports our current research approach": "supports the current research approach",
            "we do not demonstrate": "the study does not demonstrate",
            "This was especially true for recall": "This was especially true for recall",
            "less prevalent negative classes": "less prevalent positive classes",
            "The analysis also includes the file requirements for the Render deployment. I have excluded the raw public traces, generated reports, installer files, and local Word documents so as to reduce the deployment to a manageable size; I would include the source code, trained models, example data, and saved charts for an appropriate demonstration deployment.": "The analysis also includes the file requirements for the Render deployment. The raw public traces, generated reports, installer files, and local Word documents were excluded to keep the deployment manageable; the appropriate demonstration deployment includes the source code, trained models, example data, and saved charts.",
        },
    )

    # Tighten repeated Chapter Five future-work lines while preserving content.
    set_text(
        doc.paragraphs[find_first(doc, "In the future, it would be desirable to bundle the predictor")],
        (
            "Future work should package the predictor as a formal service or installable plugin with a documented API, "
            "configuration file, distribution package, and versioned model artifacts. This would make it easier for other "
            "pipelines to call the monitor without manually copying project scripts."
        ),
    )
    duplicate_starts = [
        "In the future it would be desirable to bundle the predictor as a service",
        "Another direction of future work should include adding a persistent monitoring history.",
        "Finally, it would be interesting to compare other, sequence-aware, log models",
    ]
    for start in duplicate_starts:
        try:
            idx = find_first(doc, start)
            p = doc.paragraphs[idx]
            p._element.getparent().remove(p._element)
        except ValueError:
            pass

    # Append recent references used in the revised literature review.
    add_reference(
        doc,
        "Ma, X., Li, Y., Keung, J., Yu, X., Zou, H., Yang, Z., Sarro, F., & Barr, E. T. (2024). "
        "Practitioners' expectations on log anomaly detection. arXiv. https://arxiv.org/abs/2412.01066",
    )
    add_reference(
        doc,
        "Wang, Y., Mäntylä, M. V., Nyyssölä, J., Ping, K., & Wang, L. (2024). "
        "Cross-system software log-based anomaly detection using meta-learning. arXiv. https://arxiv.org/abs/2412.15445",
    )
    add_reference(
        doc,
        "Zhang, L., Jia, T., Jia, M., Li, Y., Yang, Y., & Wu, Z. (2024). "
        "Multivariate log-based anomaly detection for distributed database. arXiv. https://arxiv.org/abs/2406.07976",
    )

    # Keep reference list roughly alphabetical after adding the new entries.
    ref_idx = para_index(doc, "REFERENCES")
    refs = [p.text for p in doc.paragraphs[ref_idx + 1:] if norm(p.text)]
    for p in list(doc.paragraphs)[ref_idx + 1:]:
        p._element.getparent().remove(p._element)
    for ref in sorted(set(refs), key=lambda s: s.lower()):
        add_reference(doc, ref)

    # Center the main title if the source had alignment drift after text replacements.
    doc.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    main()
