from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn


ROOT = Path(r"C:\Users\HP\Downloads\Final year project")
SOURCE = ROOT / "not_useful" / "documents" / "expanded_final_defence" / "Final YEAR PROPOSAL Revised.docx"
OUT = ROOT / "not_useful" / "documents" / "expanded_final_defence" / "Final YEAR PROPOSAL Revised - Chapter 3 arranged.docx"


def tag_name(el) -> str:
    return el.tag.split("}")[-1]


def text_of(el) -> str:
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


def find_body_paragraph(body, needle: str, start: int = 0) -> int:
    needle_norm = needle.strip().upper()
    for idx, el in enumerate(body):
        if idx < start or tag_name(el) != "p":
            continue
        if text_of(el).strip().upper() == needle_norm:
            return idx
    raise ValueError(f"Could not find paragraph: {needle}")


def find_body_text_startswith(body, prefix: str, start: int = 0) -> int:
    prefix_norm = prefix.strip().upper()
    for idx, el in enumerate(body):
        if idx < start or tag_name(el) != "p":
            continue
        if text_of(el).strip().upper().startswith(prefix_norm):
            return idx
    raise ValueError(f"Could not find paragraph starting with: {prefix}")


def clone_table_after_caption(body, caption_prefix: str):
    caption_idx = find_body_text_startswith(body, caption_prefix)
    for el in body[caption_idx + 1 :]:
        if tag_name(el) == "tbl":
            return deepcopy(el)
    raise ValueError(f"Could not find table after caption: {caption_prefix}")


def clone_figure_before_caption(body, caption_prefix: str):
    caption_idx = find_body_text_startswith(body, caption_prefix)
    for el in reversed(body[:caption_idx]):
        if tag_name(el) == "p" and list(el.iter(qn("w:drawing"))):
            return deepcopy(el)
    raise ValueError(f"Could not find figure before caption: {caption_prefix}")


def paragraph_element(doc: Document, text: str = "", style: str = "Body Text", page_break: bool = False):
    p = doc.add_paragraph(style=style)
    if text:
        p.add_run(text)
    if page_break:
        p.add_run().add_break(WD_BREAK.PAGE)
    el = p._element
    doc.element.body.remove(el)
    return el


def update_toc_chapter_three(doc: Document) -> None:
    """Refresh the Chapter Three TOC labels so they match the revised headings.

    Page numbers are retained as rough placeholders; the chapter headings are the
    important correction because this source document uses a static TOC rather
    than a live Word field.
    """

    toc_entries = [
        ("CHAPTER THREE", "toc 1"),
        ("RESEARCH METHODOLOGY", "toc 1"),
        ("3.1 Chapter Introduction", "toc 2"),
        ("3.2 Research Design", "toc 2"),
        ("3.2.1 Design Research Methodology", "toc 3"),
        ("3.2.2 Design Science Research Approach", "toc 3"),
        ("3.2.3 Experimental Method", "toc 3"),
        ("3.3 System and Model Design", "toc 2"),
        ("3.4 Instrumentation and Development Tools", "toc 2"),
        ("3.5 Data Collection and Sources", "toc 2"),
        ("3.6 Data Preparation and Preprocessing", "toc 2"),
        ("3.6.1 HDFS-Compatible Log Preparation", "toc 3"),
        ("3.6.2 Alibaba PAI Preparation", "toc 3"),
        ("3.6.3 Google Borg 2019 Preparation", "toc 3"),
        ("3.6.4 ETL Demonstration Preparation", "toc 3"),
        ("3.7 Feature Engineering", "toc 2"),
        ("3.8 Model Development and Validation", "toc 2"),
        ("3.9 Implementation Procedure", "toc 2"),
        ("3.9.1 External ETL Batch Process", "toc 3"),
        ("3.9.2 Live Attachment Runner", "toc 3"),
        ("3.9.3 Predictor Plugin", "toc 3"),
        ("3.9.4 Streamlit Dashboard", "toc 3"),
        ("3.9.5 Airflow DAG", "toc 3"),
        ("3.9.6 Render Deployment Preparation", "toc 3"),
        ("3.10 Risk Classification and Intervention Logic", "toc 2"),
        ("3.11 Experimental Evaluation and Method of Data Analysis", "toc 2"),
        ("3.12 Validity, Reliability and Reproducibility", "toc 2"),
        ("3.13 Ethical Considerations", "toc 2"),
        ("3.14 Chapter Summary", "toc 2"),
        ("CHAPTER FOUR", "toc 1"),
    ]

    body = doc.element.body
    ch3_idx = find_body_text_startswith(body, "CHAPTER THREE")
    ch4_idx = find_body_text_startswith(body, "CHAPTER FOUR", start=ch3_idx + 1)
    old_toc_elements = list(body)[ch3_idx : ch4_idx + 1]
    insert_at = ch3_idx
    for el in old_toc_elements:
        body.remove(el)

    # Keep conservative placeholder pages from the old TOC. The final document
    # can be refreshed in Word if exact pagination is required after formatting.
    placeholder_pages = [
        "23", "23", "23", "23", "24", "24", "25", "25", "26", "27",
        "28", "29", "29", "30", "30", "31", "32", "34", "35", "35",
        "36", "36", "37", "37", "38", "39", "40", "42", "43", "44",
    ]
    for (title, style), page in zip(toc_entries, placeholder_pages):
        el = paragraph_element(doc, f"{title}\t{page}", style=style)
        body.insert(insert_at, el)
        insert_at += 1


def main() -> None:
    doc = Document(SOURCE)
    body = doc.element.body

    # Preserve original evidence anchors before replacing the prose.
    figure_31 = clone_figure_before_caption(body, "Figure 3.1")
    figure_32 = clone_figure_before_caption(body, "Figure 3.2")
    tables = {
        "3.1": clone_table_after_caption(body, "Table 3.1"),
        "3.2": clone_table_after_caption(body, "Table 3.2"),
        "3.3": clone_table_after_caption(body, "Table 3.3"),
        "3.4": clone_table_after_caption(body, "Table 3.4"),
        "3.5": clone_table_after_caption(body, "Table 3.5"),
    }

    chapter_three_idx = find_body_paragraph(body, "CHAPTER THREE")
    research_methodology_idx = find_body_paragraph(body, "RESEARCH METHODOLOGY", chapter_three_idx + 1)
    chapter_four_idx = find_body_paragraph(body, "CHAPTER FOUR", research_methodology_idx + 1)

    # Remove old Chapter Three content, leaving the chapter title and Chapter Four intact.
    for el in list(body)[research_methodology_idx + 1 : chapter_four_idx]:
        body.remove(el)

    insert_at = research_methodology_idx + 1

    def add_para(text: str = "", style: str = "Body Text", page_break: bool = False) -> None:
        nonlocal insert_at
        el = paragraph_element(doc, text, style=style, page_break=page_break)
        body.insert(insert_at, el)
        insert_at += 1

    def add_existing(el) -> None:
        nonlocal insert_at
        body.insert(insert_at, deepcopy(el))
        insert_at += 1

    add_para("3.1 Chapter Introduction", "Heading 2")
    add_para(
        "This chapter explains the methodology used to design, implement, and evaluate the log-analysis and failure-prediction system. It moves from the research design to the system design, data sources, preprocessing, implementation, experimental evaluation, validity, and ethical considerations. The chapter is deliberately concise so that the method can be followed as a connected workflow rather than as separate descriptions of data, software, and model testing."
    )

    add_para("3.2 Research Design", "Heading 2")
    add_para(
        "The study adopted an applied design research approach supported by experimental evaluation. The design component guided the construction of a working software artifact for batch-job and ETL failure prediction, while the experimental component guided model comparison, validation, and interpretation of results. This combination was appropriate because the project was not limited to analysing historical logs; it also produced a usable monitoring prototype consisting of loaders, feature extractors, trained models, a predictor plugin, a live ETL runner, a Streamlit dashboard, report outputs, and an Airflow workflow representation."
    )

    add_para("3.2.1 Design Research Methodology", "Heading 3")
    add_para(
        "Design research was used because the project addresses a practical computing problem by creating and evaluating an artifact. The problem is the difficulty of detecting pipeline failure risk early enough for an operator to respond. The design activity therefore focused on translating this problem into a functional system that can read operational evidence, extract useful features, generate a failure-risk prediction, and present the result in a form that can support monitoring decisions."
    )

    add_para("3.2.2 Design Science Research Approach", "Heading 3")
    add_para(
        "Within design research, the study followed a Design Science Research approach. The central artifact is the failure-prediction and monitoring system. The work proceeded through problem identification, objective definition, artifact design, implementation, demonstration, evaluation, and reflection on limitations. In this project, those stages are represented by the movement from reviewed failure patterns to source-specific models, then to the predictor plugin, dashboard, reports, and orchestration evidence."
    )
    add_existing(figure_31)
    add_para("Figure 3.1: Research and prototype implementation workflow", "Body Text")
    add_para("Table 3.1: Methodological alignment with project objectives", "Body Text")
    add_existing(tables["3.1"])

    add_para("3.2.3 Experimental Method", "Heading 3")
    add_para(
        "The experimental method was used to evaluate the predictive part of the artifact. The available records were organised into labelled successful and failed outcomes, and the models were compared on held-out data. The comparison included logistic regression, decision tree, random forest, histogram gradient boosting, XGBoost, and a majority-class baseline. XGBoost was selected for the deployed prototype because it is suitable for tabular operational features and can model non-linear relationships among severity counts, retry indicators, timing values, resource indicators, and failure-related keywords. However, it was not treated as automatically superior; it was interpreted alongside the baseline and alternative classifiers."
    )

    add_para("3.3 System and Model Design", "Heading 2")
    add_para(
        "The system was designed as a source-specific machine-learning and monitoring workflow. Separate model artifacts were retained for the HDFS-compatible/ETL demonstration source, Alibaba PAI workload trace, Google Borg 2019 trace, and the earlier default HDFS-compatible path. This choice avoided forcing different schemas into a single artificial dataset and made the evidence easier to interpret."
    )
    add_para(
        "At the software level, the architecture separates the ETL process, live attachment runner, log adapter, feature extractor, prediction model, predictor plugin, dashboard, and report outputs. This separation of concerns makes the prototype easier to inspect and defend because each component has a clear responsibility. It also supports auditability: the system does not only return a prediction, but also preserves logs, features, high-risk items, resource summaries, reports, and dashboard evidence."
    )
    add_existing(figure_32)
    add_para("Figure 3.2: Live ETL attachment architecture", "Body Text")

    add_para("3.4 Instrumentation and Development Tools", "Heading 2")
    add_para(
        "The principal research instrument was the implemented software prototype. Python was used because it provides mature libraries for data manipulation, machine learning, dashboard development, process monitoring, and automation. The main libraries and tools included pandas, NumPy, scikit-learn, XGBoost, Matplotlib, Seaborn, Streamlit, psutil, and Apache Airflow. Trained models were saved in the models folder, while reproducible outputs, charts, and comparison results were saved in the outputs folder."
    )
    add_para("Table 3.3: Implementation modules used as research instruments", "Body Text")
    add_existing(tables["3.3"])
    add_para(
        "Streamlit provided the browser-based operator interface for controls, charts, logs, tables, risk summaries, and downloadable reports. Apache Airflow was included to show how the monitoring process could be represented as a task-based workflow rather than as an isolated script. Together, the instruments supported both the machine-learning experiment and the software demonstration."
    )

    add_para("3.5 Data Collection and Sources", "Heading 2")
    add_para(
        "The data for this study consisted of operational execution records showing whether computational jobs, workload tasks, or ETL pipeline runs succeeded or failed. The unit of analysis was therefore a job, task, log sequence, or observation window, not a human participant. The study used multiple sources because pipeline failure can appear differently in log-line evidence, workload traces, and controlled ETL demonstrations."
    )
    add_para(
        "The HDFS-compatible source supported log grouping, severity counting, failure-keyword extraction, and the live ETL dashboard. The Alibaba PAI GPU trace provided workload and task evidence from a large heterogeneous cluster context. The Google Borg 2019 trace provided a separate production-style workload schema. The controlled ETL demonstration supplied repeatable runtime evidence for LOW, MEDIUM, and HIGH risk scenarios in the dashboard and report outputs."
    )
    add_para("Table 3.2: Dataset sources and project role", "Body Text")
    add_existing(tables["3.2"])
    add_para(
        "A bounded purposive sampling technique was used because the public trace files are large and the project had to remain reproducible on a local development computer. The Alibaba loader produced a bounded labelled sample of 5,000 jobs, consisting of 273 failures and 4,727 successful jobs. The Google 2019 loader produced 10,000 labelled records, consisting of 2,299 failures and 7,701 successful records. The ETL demonstration generated 100 ETL records and 400 early-warning snapshots for repeatable monitoring and reporting scenarios."
    )
    add_para(
        "The bounded samples are not presented as complete representations of all rows in the public traces. They are used as reproducible evidence for the prototype and as a basis for model comparison. The larger compressed Alibaba and Google trace files were excluded from the deployable repository, while trained models, small samples, and saved comparison outputs were retained for demonstration and review."
    )

    add_para("3.6 Data Preparation and Preprocessing", "Heading 2")
    add_para(
        "Preprocessing converted each data source into modelling records while preserving the meaning of its original schema. Public traces were processed through code-based loaders rather than manual editing, and runtime ETL evidence was generated through the executable demonstration pipeline. This distinction was important because stored workload traces and live log streams require different preparation strategies."
    )
    add_para(
        "Labels were derived for supervised learning and evaluation, but terminal outcome fields were not used as predictive inputs. This prevented label leakage, where the model would appear accurate because it had access to the answer after job completion. Predictive features were limited to evidence available before or during execution, such as severity counts, retry signals, timing values, resource indicators, and failure-related keywords."
    )

    add_para("3.6.1 HDFS-Compatible Log Preparation", "Heading 3")
    add_para(
        "HDFS-compatible preparation grouped log lines by a stable block or job identifier. Each message retained its timestamp, severity level, component, stage, content, and job identity. This grouping allowed the feature extractor to summarise behaviour over a job or observation window, which is more meaningful than treating each log line as an isolated event."
    )

    add_para("3.6.2 Alibaba PAI Preparation", "Heading 3")
    add_para(
        "Alibaba PAI preparation focused on job-level and task-level workload summaries rather than natural-language log parsing. The loader extracted counts and execution indicators such as failed tasks, retries, duration-related values, and resource-related fields. The large compressed archives were kept out of the deployable repository, while the processed sample and model artifacts were retained for review."
    )

    add_para("3.6.3 Google Borg 2019 Preparation", "Heading 3")
    add_para(
        "Google Borg 2019 preparation used a local rejoined export of the Borg trace and mapped relevant workload fields into numerical modelling features. The source was treated separately from both HDFS-compatible logs and Alibaba PAI because its schema, workload representation, and failure indicators differ from those sources."
    )

    add_para("3.6.4 ETL Demonstration Preparation", "Heading 3")
    add_para(
        "The ETL demonstration used data/sample_orders.csv as a small deployable input file. The running ETL process emitted logs during extract, validate, transform, load, and report stages. Scenario controls made the demonstration repeatable by allowing normal, warning, recovering, risky, and mixed runs to be shown during defence without depending on large external trace files."
    )

    add_para("3.7 Feature Engineering", "Heading 2")
    add_para(
        "Feature engineering served as the bridge between operational evidence and machine learning. The implementation did not treat raw logs or workload records as direct model inputs; instead, it transformed them into interpretable numerical indicators that could represent failure risk."
    )
    add_para("Table 3.4: Feature engineering categories", "Body Text")
    add_existing(tables["3.4"])
    add_para(
        "For log-based evidence, the most useful features were severity counts and rates, failure keywords, retry indicators, exception terms, timing signals, and stage-level behaviour. For workload-trace evidence, the useful features were task counts, failed-task counts, resource indicators, scheduling or duration values, and workload status fields. The feature sets were kept source-specific so that a feature meaningful in one schema was not wrongly imposed on another."
    )

    add_para("3.8 Model Development and Validation", "Heading 2")
    add_para(
        "Model development followed a repeatable workflow: load labelled records, extract features, split the data, train comparative models, evaluate performance, and save the selected model artifacts. Separate artifacts were retained for the HDFS-compatible/ETL demonstration, Alibaba PAI, Google 2019, and legacy/default HDFS-compatible paths."
    )
    add_para(
        "The implemented XGBoost configuration used 300 estimators, a maximum tree depth of 5, a learning rate of 0.05, subsampling of 0.9, column subsampling of 0.9, and log-loss as the evaluation metric. Held-out metrics and comparison outputs were saved as project evidence. Five-fold cross-validation was also used, and StratifiedGroupKFold was applied where repeated snapshots or grouped observations made ordinary random splitting unsafe."
    )
    add_para("Table 3.5: Stored model artifacts", "Body Text")
    add_existing(tables["3.5"])
    add_para(
        "A majority-class baseline was included because failure-prediction datasets are often imbalanced. Accuracy alone can therefore be misleading if most jobs succeed. The validation process considered precision, recall, F1-score, ROC-AUC, PR-AUC, MCC, and confusion-matrix counts so that the failure class could be interpreted more carefully."
    )

    add_para("3.9 Implementation Procedure", "Heading 2")
    add_para(
        "Implementation proceeded from offline modelling to a live monitoring workflow. The first stage produced loaders, feature extraction, model comparison, and saved artifacts. The next stage connected the prediction logic to an executable ETL process, then added the predictor plugin, dashboard, report generator, Airflow DAG, and Render deployment preparation."
    )

    add_para("3.9.1 External ETL Batch Process", "Heading 3")
    add_para(
        "The external ETL batch process followed the familiar extract, validate, transform, load, and report sequence. Each phase emitted progress messages, allowing the monitor to observe normal execution, warnings, recovery behaviour, and risky behaviour. Scenario controls made the process repeatable for demonstration."
    )

    add_para("3.9.2 Live Attachment Runner", "Heading 3")
    add_para(
        "The live attachment runner executed the external command, captured stdout and stderr, normalised incoming lines, and forwarded structured records to the predictor plugin. Where available, it also captured CPU, memory, and disk readings through psutil, giving the dashboard additional execution context."
    )

    add_para("3.9.3 Predictor Plugin", "Heading 3")
    add_para(
        "The predictor plugin made the monitoring logic reusable. It ingested standardised records, updated feature summaries, saved alerts, recognised HIGH-risk events, maintained review state, and produced report payloads in HTML and JSON. This prevented prediction logic from being scattered across the ETL script, dashboard, and report generator."
    )

    add_para("3.9.4 Streamlit Dashboard", "Heading 3")
    add_para(
        "The Streamlit dashboard served as the operator-facing and examiner-facing layer of the prototype. It displayed controls, stage indicators, logs, metrics, risk cards, model comparison charts, system maps, and downloadable report outputs. Its role was to make the monitoring workflow visible rather than merely display a final model score."
    )

    add_para("3.9.5 Airflow DAG", "Heading 3")
    add_para(
        "The Airflow DAG represented the workflow as tasks: prepare_run, run_pipeline_with_predictor, evaluate_intervention, and publish_report. It was included as orchestration evidence showing where the predictor could sit inside a task-based pipeline. It was not presented as part of the Render web service."
    )

    add_para("3.9.6 Render Deployment Preparation", "Heading 3")
    add_para(
        "Render deployment preparation focused on the Streamlit dashboard because it was the part of the system most suitable for a lightweight hosted demonstration. The repository retained render.yaml, .python-version, requirements.txt, trained models, small sample data, and output graphs. Larger raw traces, local report drafts, and generated heavyweight outputs were kept outside the deployable repository."
    )

    add_para("3.10 Risk Classification and Intervention Logic", "Heading 2")
    add_para(
        "The system translated model probabilities into three operator-friendly risk categories. A probability greater than or equal to 0.70 was classified as HIGH, a probability from 0.40 to below 0.70 was classified as MEDIUM, and a probability below 0.40 was classified as LOW. These thresholds made the dashboard and report easier to interpret during demonstration."
    )
    add_para(
        "The implementation distinguished between prediction and intervention. HIGH risk did not automatically mean that the pipeline had to be stopped in every context. In advisory mode, the system warns the operator; in review mode, the operator decides whether to continue; and in auto-stop mode, the system can halt execution based on configuration. This supports decision-making without overstating the model as a complete replacement for human judgement."
    )
    add_para(
        "The system also preserved the distinction between raw model probability and demonstration calibration. Where a presentational value was used to make risk movement visible during defence, raw_model_failure_probability was retained so that evaluation evidence remained separate from dashboard presentation behaviour."
    )

    add_para("3.11 Experimental Evaluation and Method of Data Analysis", "Heading 2")
    add_para(
        "Data analysis was conducted at two levels. At the model level, saved metrics, comparison files, confusion matrices, and charts were interpreted to assess predictive performance. At the system level, executable evidence from the live runner, ETL process, predictor plugin, dashboard, report generator, and Airflow DAG was examined to confirm that the artifact functioned as a monitoring workflow."
    )
    add_para(
        "Accuracy was reported but not treated as sufficient because successful jobs are usually more common than failed jobs. Precision, recall, F1-score, PR-AUC, ROC-AUC, MCC, and confusion-matrix counts provided a more balanced interpretation of the failure class. Chapter Four uses these outputs to connect the experimental results with the implemented system behaviour."
    )

    add_para("3.12 Validity, Reliability and Reproducibility", "Heading 2")
    add_para(
        "Internal validity was supported by avoiding label leakage, preserving source-specific schemas, and comparing the selected model against alternative classifiers and a majority baseline. Construct validity was supported by using interpretable operational indicators such as severity counts, retry signals, exception keywords, timing information, resource indicators, and failed-task counts."
    )
    add_para(
        "Reliability was supported by the modular structure of the implementation. Data loading, feature creation, model training, model comparison, live running, prediction, reporting, and dashboard rendering were separated into inspectable components. Reproducibility was supported by saving model artifacts, metrics, graphs, sample data, configuration files, and report outputs."
    )
    add_para(
        "The study also had clear boundaries. The ETL demonstration was controlled, the Alibaba and Google samples were bounded, the dashboard could include presentation calibration, and the Airflow DAG was a proof of orchestration design rather than a hosted Render service. The project therefore claims a practicable prototype and workflow, not a fully production-validated failure-prediction service."
    )

    add_para("3.13 Ethical Considerations", "Heading 2")
    add_para(
        "The study used public traces and synthetic demonstration data, not data collected from human participants. Nevertheless, operational logs can contain sensitive business information, credentials, or personally identifiable information in real deployments. A production version would therefore require log masking, access control, report-view restrictions, and procedures for handling sensitive data."
    )
    add_para(
        "The ethical position of the project is also reflected in how the model is presented. The prototype supports human decision-making but does not claim to replace operational judgement. Its limitations are stated openly so that demonstration behaviour, model evidence, and production readiness are not confused."
    )

    add_para("3.14 Chapter Summary", "Heading 2")
    add_para(
        "This chapter presented the research design, Design Science Research approach, experimental method, system and model design, data sources, preprocessing, feature engineering, implementation procedure, risk logic, evaluation method, validity, reproducibility, and ethical considerations. The methodology links the developed artifact to the evidence used to evaluate it, preparing the ground for the results and analysis presented in Chapter Four.",
        page_break=True,
    )

    update_toc_chapter_three(doc)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
